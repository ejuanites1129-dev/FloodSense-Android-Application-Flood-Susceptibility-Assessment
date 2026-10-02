"""Bounded, reviewed CSV/XLSX staging. No fetches or automatic publication."""

import csv
import io
import ntpath
from collections import Counter
from decimal import Decimal, InvalidOperation
from xml.etree.ElementTree import ParseError
from zipfile import BadZipFile, ZipFile

from defusedxml.common import DefusedXmlException
from django.contrib.admin.models import ADDITION, CHANGE, LogEntry
from django.contrib.gis.geos import Point
from django.core.exceptions import ValidationError
from django.db import transaction
from django.utils import timezone
from geography.models import GeographicArea
from provenance.models import DataSource, PublicationStatus

from .models import CenterImportBatch, EvacuationCenter

HEADERS = ("name", "address", "area_code", "latitude", "longitude", "limitations")
MAX_BYTES = 5 * 1024 * 1024
MAX_ROWS = 1000


def read_upload(upload):
    if upload.size > MAX_BYTES:
        raise ValidationError("Upload is limited to 5 MB and 1,000 data rows.")
    raw = upload.read(MAX_BYTES + 1)
    if len(raw) > MAX_BYTES:
        raise ValidationError("Upload exceeds 5 MB.")
    filename = ntpath.basename(upload.name)
    try:
        if filename.lower().endswith(".csv"):
            table = []
            for row in csv.reader(io.StringIO(raw.decode("utf-8-sig")), strict=True):
                if len(table) > MAX_ROWS:
                    raise ValidationError("Upload exceeds the 1,000 data row limit.")
                table.append(row)
        elif filename.lower().endswith(".xlsx"):
            from openpyxl import load_workbook

            with ZipFile(io.BytesIO(raw)) as archive:
                entries = archive.infolist()
                if len(entries) > 1000 or sum(e.file_size for e in entries) > 20 * 1024 * 1024:
                    raise ValidationError("Workbook expands beyond the safe processing limit.")
                if any(
                    "vbaProject" in e.filename or "externalLinks/" in e.filename for e in entries
                ):
                    raise ValidationError("Macros and external workbook links are not accepted.")
            workbook = load_workbook(
                io.BytesIO(raw), read_only=True, data_only=False, keep_links=False
            )
            try:
                if len(workbook.sheetnames) != 1:
                    raise ValidationError("Use one worksheet with the exact template headers.")
                sheet = workbook.active
                # Ignore untrusted worksheet dimension metadata; bound actual streamed cells.
                sheet.reset_dimensions()
                table = []
                for cells in sheet.iter_rows():
                    if len(table) > MAX_ROWS or len(cells) > len(HEADERS):
                        raise ValidationError("Workbook exceeds the row or column limit.")
                    table.append(["" if cell.value is None else str(cell.value) for cell in cells])
            finally:
                workbook.close()
        else:
            raise ValidationError(
                "Use UTF-8 CSV or .xlsx. Legacy .xls and macro workbooks are not accepted."
            )
    except (
        UnicodeError,
        csv.Error,
        BadZipFile,
        KeyError,
        OSError,
        ValueError,
        DefusedXmlException,
        ParseError,
    ) as error:
        raise ValidationError(
            "The file could not be read. Use the exact CSV/Excel template."
        ) from error
    if not table or tuple(table[0]) != HEADERS:
        raise ValidationError(
            "Headers must exactly match the downloaded template, in the same order."
        )
    if len(table) < 2 or len(table) - 1 > MAX_ROWS:
        raise ValidationError(
            "Supply 1 through 1,000 data rows; blank rows are reported, not discarded."
        )
    if any(len(str(cell)) > 10000 for row in table for cell in row):
        raise ValidationError("A cell exceeds the safe text length limit.")
    return filename[:180], table[1:]


def _coordinate_key(row):
    if len(row) < 5:
        return None
    try:
        values = (Decimal(row[3]), Decimal(row[4]))
        return values if all(v.is_finite() for v in values) else None
    except InvalidOperation:
        return None


def validate_rows(table, source, shared_limitations=""):
    """Validate EVERY row without writes. Duplicate candidates cannot auto-import."""
    areas = {a.code: a for a in GeographicArea.objects.select_related("source")}
    existing = list(EvacuationCenter.objects.values("name", "latitude", "longitude"))
    names = Counter(row[0].strip().casefold() for row in table if row)
    coordinates = Counter(_coordinate_key(row) for row in table)
    results = []
    for number, raw in enumerate(table, start=2):
        errors, warnings = [], []
        values = dict(zip(HEADERS, raw, strict=False))
        if len(raw) != len(HEADERS):
            errors.append("Wrong column count; no fields were silently discarded.")
        if any(str(v).lstrip().startswith("=") for v in raw):
            errors.append("Formulas are not accepted; supply literal documented values.")
        if any(any(ord(c) < 32 and c not in "\n\r\t" for c in v) for v in raw):
            errors.append("Unsupported control characters; supply safe literal text.")
        area = areas.get(values.get("area_code", ""))
        if values.get("area_code") and area is None:
            errors.append(
                "Unknown area_code; use an existing supported area identity, not a guessed name."
            )
        elif area is None:
            warnings.append("No geographic identity assigned; excluded from resident results.")
        if names[values.get("name", "").strip().casefold()] > 1:
            errors.append(
                "Possible duplicate normalized name within this batch; resolve before import."
            )
        center = EvacuationCenter(
            name=values.get("name", ""),
            address=values.get("address", ""),
            latitude=values.get("latitude", ""),
            longitude=values.get("longitude", ""),
            geographic_area=area,
            source=source,
            limitations=values.get("limitations", "") or shared_limitations,
        )
        try:
            center.full_clean()
        except ValidationError as error:
            errors.extend(error.messages)
        else:
            if coordinates[_coordinate_key(raw)] > 1 or any(
                e["name"].strip().casefold() == center.name.strip().casefold()
                or (e["latitude"] == center.latitude and e["longitude"] == center.longitude)
                for e in existing
            ):
                errors.append(
                    "Possible duplicate name/coordinates; resolve or enter "
                    "a separately reviewed draft."
                )
            if area:
                if not area.geometry or area.geometry.empty or not area.geometry.valid:
                    errors.append("Associated geographic geometry is missing or invalid.")
                elif not area.geometry.covers(
                    Point(float(center.longitude), float(center.latitude), srid=4326)
                ):
                    errors.append(
                        "Coordinate lies outside the associated area's recorded geometry."
                    )
                if not area.is_enabled:
                    warnings.append("Area disabled; not resident-eligible.")
        results.append(
            {
                "number": number,
                "values": values,
                "errors": errors,
                "warnings": warnings,
                "area_revision": area.updated_at.isoformat() if area else None,
            }
        )
    return results


@transaction.atomic
def stage_batch(*, upload, source, actor, shared_limitations=""):
    from core.record_workflow import require_permission

    require_permission(actor, "evacuation.add_evacuationcenter")
    source = DataSource.objects.select_for_update().get(pk=source.pk)
    if source.status in {
        PublicationStatus.RESTRICTED,
        PublicationStatus.RETIRED,
        PublicationStatus.DEMONSTRATION,
    }:
        raise ValidationError("Use an authorized pending or approved non-demonstration source.")
    filename, table = read_upload(upload)
    rows = validate_rows(table, source, shared_limitations)
    # Staging metadata is written only AFTER all rows have been validated.
    valid = sum(not row["errors"] for row in rows)
    batch = CenterImportBatch.objects.create(
        source=source,
        actor=actor,
        filename=filename,
        rows=rows,
        shared_limitations=shared_limitations,
        source_revision=source.updated_at,
        row_count=len(rows),
        valid_count=valid,
        invalid_count=len(rows) - valid,
    )
    LogEntry.objects.log_actions(
        user_id=actor.pk,
        queryset=[batch],
        action_flag=ADDITION,
        change_message="Staged center batch; no center records imported.",
    )
    return batch


@transaction.atomic
def import_batch(*, batch_id, actor):
    from core.record_workflow import center_write_lock, require_permission

    from evacuation.workflow import log_center_action

    require_permission(actor, "evacuation.add_evacuationcenter")
    center_write_lock()
    batch = CenterImportBatch.objects.select_for_update().get(pk=batch_id)
    source = DataSource.objects.select_for_update().get(pk=batch.source_id)
    if batch.imported_at:
        raise ValidationError("This batch was already imported; no records were created again.")
    if source.updated_at != batch.source_revision:
        raise ValidationError(
            "Source metadata changed after preview. Upload again and review the new validation."
        )
    # Keep every referenced geometry fixed through validation and insertion.
    list(
        GeographicArea.objects.select_for_update()
        .filter(code__in=[row["values"].get("area_code", "") for row in batch.rows])
        .order_by("pk")
    )
    valid = [row for row in batch.rows if not row["errors"]]
    if not valid:
        raise ValidationError("There are no valid rows to import.")
    fresh = validate_rows(
        [[r["values"].get(h, "") for h in HEADERS] for r in valid], source, batch.shared_limitations
    )
    if any(
        r["errors"] or r["area_revision"] != old["area_revision"]
        for r, old in zip(fresh, valid, strict=True)
    ):
        raise ValidationError(
            "Rows or geographic context changed after preview. Nothing imported; upload again."
        )
    for row in valid:
        values = row["values"]
        area = GeographicArea.objects.filter(code=values["area_code"]).first()
        center = EvacuationCenter(
            name=values["name"],
            address=values["address"],
            latitude=values["latitude"],
            longitude=values["longitude"],
            geographic_area=area,
            source=source,
            limitations=values["limitations"] or batch.shared_limitations,
        )
        center.full_clean()
        center.save()
        row["center_id"] = center.pk
        log_center_action(
            actor=actor,
            center=center,
            action_flag=ADDITION,
            message=f"Created pending draft from reviewed batch {batch.public_id}.",
        )
    batch.imported_count = len(valid)
    batch.imported_at = timezone.now()
    batch.imported_by = actor
    batch.save(update_fields=("rows", "imported_count", "imported_at", "imported_by"))
    LogEntry.objects.log_actions(
        user_id=actor.pk,
        queryset=[batch],
        action_flag=CHANGE,
        change_message=(
            f"Imported {batch.imported_count} reviewed draft rows; "
            f"{batch.invalid_count} invalid rows explicitly excluded."
        ),
    )
    return batch
